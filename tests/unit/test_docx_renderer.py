from docx import Document

from send_email_studio.infrastructure.docx_renderer import render_docx


def test_docx_renderer_keeps_table_and_basic_formatting(tmp_path):
    path = tmp_path / "template.docx"
    document = Document()
    paragraph = document.add_paragraph()
    run = paragraph.add_run("标题")
    run.bold = True
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "姓名"
    table.cell(0, 1).text = "{姓名}"
    document.save(path)
    html = render_docx(path)
    assert "<table" in html and "姓名" in html and "font-weight:700" in html
