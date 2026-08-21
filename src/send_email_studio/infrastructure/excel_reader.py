from pathlib import Path
from numbers import Real
import re
from typing import Any

import pandas as pd
from openpyxl import load_workbook


_PERCENT_DECIMALS_RE = re.compile(r"\.([0#]+)%")


def _display_value(value: Any, number_format: str) -> Any:
    """将 Excel 百分比单元格还原为用户看到的文本，而非底层小数。"""
    if value is None:
        return ""
    if isinstance(value, Real) and not isinstance(value, bool) and "%" in number_format:
        section = number_format.split(";", 1)[0]
        match = _PERCENT_DECIMALS_RE.search(section)
        decimals = len(match.group(1)) if match else 0
        return f"{float(value) * 100:.{decimals}f}%"
    return value


def read_excel_rows(path: str | Path) -> tuple[list[str], list[dict[str, Any]]]:
    file_path = Path(path)
    if file_path.suffix.lower() not in {".xlsx", ".xlsm", ".xls"}:
        raise ValueError("仅支持 .xlsx、.xlsm 或 .xls 文件")
    if not file_path.is_file():
        raise FileNotFoundError(f"Excel 文件不存在：{file_path}")
    frame = pd.read_excel(file_path, dtype=object).fillna("")
    columns = [str(c) for c in frame.columns]
    rows = [{str(k): value for k, value in row.items()} for row in frame.to_dict(orient="records")]
    if file_path.suffix.lower() in {".xlsx", ".xlsm"} and rows:
        workbook = load_workbook(file_path, data_only=True, read_only=True)
        try:
            sheet = workbook.active
            for row_index, cells in enumerate(sheet.iter_rows(min_row=2, max_col=len(columns)), start=0):
                if row_index >= len(rows):
                    break
                for column_index, cell in enumerate(cells):
                    if "%" in cell.number_format:
                        rows[row_index][columns[column_index]] = _display_value(
                            cell.value, cell.number_format
                        )
        finally:
            workbook.close()
    return columns, rows
